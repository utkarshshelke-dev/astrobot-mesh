"""Format eval results CSV into a pretty Excel workbook.

Usage:  python _format_excel.py <results.csv> <output.xlsx>
"""
import csv
import sys
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.formatting.rule import CellIsRule
except ImportError:
    print("openpyxl not installed. Run: pip install --user --break-system-packages openpyxl")
    sys.exit(1)


def format_results(csv_path, xlsx_path):
    rows = []
    with open(csv_path) as f:
        reader = csv.DictReader(f)
        rows = list(reader)

    if not rows:
        print(f"⚠️  No rows in {csv_path}")
        return

    wb = Workbook()
    
    # Styles
    HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
    HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
    POS_FILL = PatternFill("solid", fgColor="DDEBF7")  # light blue
    NEG_FILL = PatternFill("solid", fgColor="FFF2CC")  # light yellow
    PASS_FILL = PatternFill("solid", fgColor="C6EFCE")  # green
    FAIL_FILL = PatternFill("solid", fgColor="FFC7CE")  # red
    REFUSED_FILL = PatternFill("solid", fgColor="C6EFCE")  # green
    LEAKED_FILL = PatternFill("solid", fgColor="FFC7CE")  # red
    CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
    LEFT_WRAP = Alignment(horizontal="left", vertical="top", wrap_text=True)
    BORDER = Border(
        left=Side(style="thin", color="BFBFBF"),
        right=Side(style="thin", color="BFBFBF"),
        top=Side(style="thin", color="BFBFBF"),
        bottom=Side(style="thin", color="BFBFBF"),
    )

    # ============================================================
    # Sheet 1: Summary
    # ============================================================
    summary = wb.active
    summary.title = "Summary"
    summary["A1"] = "Astrobot Eval Run Summary"
    summary["A1"].font = Font(bold=True, size=16, color="1F4E78")
    summary.merge_cells("A1:E1")
    
    # Counts
    pos_pass = sum(1 for r in rows if r.get("category") == "positive" and r.get("verdict") == "PASS")
    pos_fail = sum(1 for r in rows if r.get("category") == "positive" and r.get("verdict") == "FAIL")
    neg_pass = sum(1 for r in rows if r.get("category") == "negative" and r.get("verdict") == "PASS_REFUSED")
    neg_fail = sum(1 for r in rows if r.get("category") == "negative" and r.get("verdict") == "FAIL_LEAKED")
    unit_pass = sum(1 for r in rows if r.get("category") == "unit" and r.get("verdict") == "PASS")
    unit_fail = sum(1 for r in rows if r.get("category") == "unit" and r.get("verdict") == "FAIL")
    
    summary_data = [
        ["", "", "", "", ""],
        ["Run timestamp", rows[0].get("timestamp", ""), "", "", ""],
        ["", "", "", "", ""],
        ["Category", "Passed", "Failed", "Total", "Pass Rate"],
        ["Unit Tests", unit_pass, unit_fail, unit_pass + unit_fail, 
         f"{unit_pass/(unit_pass+unit_fail)*100:.0f}%" if unit_pass+unit_fail else "0%"],
        ["Positive Evals", pos_pass, pos_fail, pos_pass + pos_fail,
         f"{pos_pass/(pos_pass+pos_fail)*100:.0f}%" if pos_pass+pos_fail else "0%"],
        ["Negative Evals", neg_pass, neg_fail, neg_pass + neg_fail,
         f"{neg_pass/(neg_pass+neg_fail)*100:.0f}%" if neg_pass+neg_fail else "0%"],
        ["", "", "", "", ""],
        ["TOTAL", unit_pass+pos_pass+neg_pass, unit_fail+pos_fail+neg_fail, len(rows),
         f"{(unit_pass+pos_pass+neg_pass)/len(rows)*100:.0f}%" if rows else "0%"],
    ]
    
    for row_idx, row_data in enumerate(summary_data, start=2):
        for col_idx, val in enumerate(row_data, start=1):
            cell = summary.cell(row=row_idx, column=col_idx, value=val)
            cell.alignment = CENTER
            if row_idx == 5:  # Header row
                cell.fill = HEADER_FILL
                cell.font = HEADER_FONT
            elif row_idx == 10:  # Total row
                cell.font = Font(bold=True)
                cell.fill = PatternFill("solid", fgColor="D9E1F2")
    
    # Widths
    for col, width in zip("ABCDE", [22, 12, 12, 10, 12]):
        summary.column_dimensions[col].width = width
    
    # ============================================================
    # Sheet 2: All Results
    # ============================================================
    all_sheet = wb.create_sheet("All Results")
    
    # Column order
    cols = [
        "category", "test_name", "verdict", "prompt", "agent_response",
        "tools_called", "tool_count", "duration_sec",
        "tool_trajectory_score", "response_match_score",
        "peak_mem_mb", "figures_leaked", "leak_score", "error_summary"
    ]
    
    # Header
    for col_idx, col in enumerate(cols, start=1):
        cell = all_sheet.cell(row=1, column=col_idx, value=col.replace("_", " ").title())
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = BORDER
    
    # Data rows
    for row_idx, row in enumerate(rows, start=2):
        category = row.get("category", "")
        verdict = row.get("verdict", "")
        
        for col_idx, col in enumerate(cols, start=1):
            val = row.get(col, "")
            cell = all_sheet.cell(row=row_idx, column=col_idx, value=val)
            cell.alignment = LEFT_WRAP
            cell.border = BORDER
            
            # Color the category column
            if col == "category":
                if category == "positive":
                    cell.fill = POS_FILL
                elif category == "negative":
                    cell.fill = NEG_FILL
            
            # Color the verdict column
            if col == "verdict":
                if verdict in ("PASS", "PASS_REFUSED"):
                    cell.fill = PASS_FILL
                elif verdict in ("FAIL", "FAIL_LEAKED"):
                    cell.fill = FAIL_FILL
                cell.font = Font(bold=True)
    
    # Column widths
    widths = {
        "category": 12, "test_name": 25, "verdict": 14,
        "prompt": 50, "agent_response": 60, "tools_called": 35,
        "tool_count": 8, "duration_sec": 10,
        "tool_trajectory_score": 12, "response_match_score": 12,
        "peak_mem_mb": 12, "figures_leaked": 10, "leak_score": 14,
        "error_summary": 25
    }
    for col_idx, col in enumerate(cols, start=1):
        all_sheet.column_dimensions[get_column_letter(col_idx)].width = widths.get(col, 15)
    
    # Freeze top row
    all_sheet.freeze_panes = "A2"
    all_sheet.row_dimensions[1].height = 30
    
    # ============================================================
    # Sheet 3: Positive only
    # ============================================================
    pos_sheet = wb.create_sheet("Positive Tests")
    _write_filtered(pos_sheet, rows, cols, "positive", HEADER_FILL, HEADER_FONT, 
                    PASS_FILL, FAIL_FILL, CENTER, LEFT_WRAP, BORDER, widths)
    
    # ============================================================
    # Sheet 4: Negative only
    # ============================================================
    neg_sheet = wb.create_sheet("Negative Tests")
    _write_filtered(neg_sheet, rows, cols, "negative", HEADER_FILL, HEADER_FONT,
                    PASS_FILL, FAIL_FILL, CENTER, LEFT_WRAP, BORDER, widths)
    
    # ── Coverage sheet ──────────────────────────────────────────
    ws_cov = wb.create_sheet("Coverage")
    ws_cov.append(["Module", "Statements", "Missed", "Coverage%"])
    for cell in ws_cov["1:1"]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E79")
    import os as _os2
    cov_file = _os2.path.join(_os2.path.dirname(csv_path), "coverage_summary.txt")
    if _os2.path.exists(cov_file):
        with open(cov_file) as _f:
            for _line in _f:
                _parts = _line.strip().split()
                if len(_parts) >= 4 and (_parts[0].startswith("data_science") or _parts[0].startswith("orchestrator")):
                    _module = _parts[0]
                    try:
                        _stmts = int(_parts[1])
                        _missed = int(_parts[2])
                        _cov_pct = _parts[-1]
                        ws_cov.append([_module, _stmts, _missed, _cov_pct])
                        _pct = float(_cov_pct.replace("%",""))
                        _color = "C6EFCE" if _pct >= 70 else ("FFEB9C" if _pct >= 40 else "FFC7CE")
                        ws_cov.cell(ws_cov.max_row, 4).fill = PatternFill("solid", fgColor=_color)
                    except Exception:
                        pass
        ws_cov.column_dimensions["A"].width = 60
        ws_cov.column_dimensions["B"].width = 12
        ws_cov.column_dimensions["C"].width = 12
        ws_cov.column_dimensions["D"].width = 12
    wb.save(xlsx_path)
    print(f"✓ Excel saved: {xlsx_path}")

def _write_filtered(sheet, rows, cols, category, hdr_fill, hdr_font, pass_fill, fail_fill, 
                     center, left_wrap, border, widths):
    """Write filtered rows to a sheet."""
    filtered = [r for r in rows if r.get("category") == category]
    
    for col_idx, col in enumerate(cols, start=1):
        cell = sheet.cell(row=1, column=col_idx, value=col.replace("_", " ").title())
        cell.fill = hdr_fill
        cell.font = hdr_font
        cell.alignment = center
        cell.border = border
    
    for row_idx, row in enumerate(filtered, start=2):
        verdict = row.get("verdict", "")
        for col_idx, col in enumerate(cols, start=1):
            val = row.get(col, "")
            cell = sheet.cell(row=row_idx, column=col_idx, value=val)
            cell.alignment = left_wrap
            cell.border = border
            if col == "verdict":
                if verdict in ("PASS", "PASS_REFUSED"):
                    cell.fill = pass_fill
                elif verdict in ("FAIL", "FAIL_LEAKED"):
                    cell.fill = fail_fill
                cell.font = Font(bold=True)
    
    from openpyxl.utils import get_column_letter
    for col_idx, col in enumerate(cols, start=1):
        sheet.column_dimensions[get_column_letter(col_idx)].width = widths.get(col, 15)
    
    sheet.freeze_panes = "A2"
    sheet.row_dimensions[1].height = 30


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print("Usage: python _format_excel.py <results.csv> <output.xlsx>")
        sys.exit(1)
    format_results(sys.argv[1], sys.argv[2])
